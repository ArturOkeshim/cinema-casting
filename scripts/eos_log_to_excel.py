#!/usr/bin/env python3
"""
Преобразует logs/eos-debug.jsonl в Excel для анализа EOS.

Использование:
  pip install openpyxl
  python scripts/eos_log_to_excel.py logs/eos-debug-server.jsonl
  python scripts/eos_log_to_excel.py logs/eos-debug-server.jsonl -o logs/report.xlsx
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Нужен пакет openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


TURN_END_COLUMNS = [
    ("ts", "Время (UTC)"),
    ("sessionId", "Сессия"),
    ("role", "Роль"),
    ("finishReason", "Завершение"),
    ("seqIdx", "Шаг"),
    ("actorTurnIndex", "Реплика №"),
    ("speakableText", "Эталон"),
    ("hypothesisRaw", "Распознано RAW"),
    ("hypothesisTrimmed", "Распознано TRIM"),
    ("trimWordsSkipped", "слов отрезано"),
    ("trimApplied", "trim?"),
    ("failedGates", "Не прошло (trim)"),
    ("tailTrim", "tail trim"),
    ("scoreTrim", "score trim"),
    ("lenTrim", "len trim"),
    ("tailRaw", "tail raw"),
    ("scoreRaw", "score raw"),
    ("lenRaw", "len raw"),
    ("tail_margin", "не хватило tail"),
    ("score_margin", "не хватило score"),
    ("partialWouldPass", "partial ок? (trim)"),
    ("partialCount", "partials"),
    ("finalCount", "finals"),
    ("turnDurationSec", "сек реплики"),
    ("smResumeDelayMs", "SM delay ms"),
    ("smError", "ошибка SM"),
    ("minLenRatio", "порог len"),
    ("scoreThreshold", "порог score"),
    ("minTailScore", "порог tail"),
]

FINALS_COLUMNS = [
    ("ts", "Время"),
    ("sessionId", "Сессия"),
    ("seqIdx", "Шаг"),
    ("finalIndex", "final №"),
    ("tMs", "мс от старта"),
    ("segmentText", "сегмент SM"),
    ("hypothesisRaw", "гипотеза RAW"),
    ("hypothesisTrimmed", "гипотеза TRIM"),
    ("tailTrim", "tail trim"),
    ("scoreTrim", "score trim"),
    ("passedTrim", "passed trim"),
    ("failedGatesTrim", "не прошло"),
]


def load_jsonl(path: Path) -> list[dict]:
    rows: list[dict] = []
    with path.open(encoding="utf-8") as f:
        for line_no, line in enumerate(f, 1):
            line = line.strip()
            if not line:
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError as e:
                print(f"Пропуск строки {line_no}: {e}", file=sys.stderr)
    return rows


def session_roles(events: list[dict]) -> dict[str, str]:
    out: dict[str, str] = {}
    for ev in events:
        if ev.get("event") == "rehearsal_start":
            sid = ev.get("sessionId")
            if sid:
                out[sid] = ev.get("role") or ""
    return out


def margins_row(gate_margins: dict | None) -> dict[str, float | None]:
    gm = gate_margins or {}
    return {
        "tail_margin": gm.get("tail"),
        "score_margin": gm.get("score"),
        "len_margin": gm.get("lenRatio"),
    }


def flatten_turn_end(ev: dict, roles: dict[str, str]) -> dict:
    m_trim = ev.get("metricsTrimmed") or ev.get("metricsFinal") or {}
    m_raw = ev.get("metricsRaw") or {}
    th = ev.get("thresholds") or {}
    gm = margins_row(ev.get("gateMarginsTrimmed") or ev.get("gateMarginsFinal"))
    failed = ev.get("failedGatesTrimmed") or ev.get("failedGatesFinal") or []
    return {
        "ts": ev.get("ts", ""),
        "sessionId": ev.get("sessionId", ""),
        "role": ev.get("role") or roles.get(ev.get("sessionId", ""), ""),
        "finishReason": ev.get("finishReason", ""),
        "seqIdx": ev.get("seqIdx"),
        "actorTurnIndex": ev.get("actorTurnIndex"),
        "speakableText": ev.get("speakableText", ""),
        "hypothesisRaw": ev.get("hypothesisRaw") or ev.get("hypothesisFinal", ""),
        "hypothesisTrimmed": ev.get("hypothesisTrimmed") or ev.get("hypothesisFinal", ""),
        "trimWordsSkipped": ev.get("trimWordsSkipped", 0),
        "trimApplied": ev.get("trimApplied", False),
        "failedGates": ", ".join(failed),
        "tailTrim": m_trim.get("tail"),
        "scoreTrim": m_trim.get("score"),
        "lenTrim": m_trim.get("lenRatio"),
        "tailRaw": m_raw.get("tail"),
        "scoreRaw": m_raw.get("score"),
        "lenRaw": m_raw.get("lenRatio"),
        **gm,
        "partialWouldPass": ev.get("partialWouldPassTrimmed") or ev.get("partialWouldPass"),
        "partialCount": ev.get("partialCount"),
        "finalCount": ev.get("finalCount"),
        "turnDurationSec": round((ev.get("turnDurationMs") or 0) / 1000, 1),
        "smResumeDelayMs": ev.get("smResumeDelayMs", 0),
        "smError": ev.get("smError") or "",
        "minLenRatio": th.get("minLenRatio"),
        "scoreThreshold": th.get("scoreThreshold"),
        "minTailScore": th.get("minTailScore"),
    }


def flatten_finals_from_turn(ev: dict) -> list[dict]:
    rows = []
    for item in ev.get("timeline") or []:
        if item.get("kind") != "final":
            continue
        m_trim = item.get("metricsTrimmed") or {}
        rows.append({
            "ts": ev.get("ts", ""),
            "sessionId": ev.get("sessionId", ""),
            "seqIdx": ev.get("seqIdx"),
            "finalIndex": item.get("finalIndex"),
            "tMs": item.get("tMs"),
            "segmentText": item.get("segmentText", ""),
            "hypothesisRaw": item.get("hypothesisRaw", ""),
            "hypothesisTrimmed": item.get("hypothesisTrimmed", ""),
            "tailTrim": m_trim.get("tail"),
            "scoreTrim": m_trim.get("score"),
            "passedTrim": item.get("passedTrimmed"),
            "failedGatesTrim": ", ".join(item.get("failedGatesTrimmed") or []),
        })
    return rows


def write_sheet(ws, columns: list[tuple[str, str]], data: list[dict]) -> None:
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    keys = [c[0] for c in columns]
    headers = [c[1] for c in columns]

    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, row in enumerate(data, 2):
        for col_idx, key in enumerate(keys, 1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(key))

    for col_idx in range(1, len(columns) + 1):
        letter = get_column_letter(col_idx)
        max_len = len(str(headers[col_idx - 1]))
        for row in data[:200]:
            val = row.get(keys[col_idx - 1])
            if val is not None:
                max_len = max(max_len, min(len(str(val)), 80))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 60)

    ws.freeze_panes = "A2"
    if data:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(data) + 1}"


def build_summary(events: list[dict], roles: dict[str, str]) -> list[dict]:
    counter = Counter(ev.get("event", "?") for ev in events)
    turn = Counter(
        ev.get("finishReason", "?")
        for ev in events
        if ev.get("event") == "turn_end"
    )
    manual_failed = Counter()
    for ev in events:
        if ev.get("event") != "turn_end" or ev.get("finishReason") != "manual":
            continue
        for g in ev.get("failedGatesTrimmed") or ev.get("failedGatesFinal") or []:
            manual_failed[g] += 1

    rows = [
        {"Показатель": "Всего строк в логе", "Значение": len(events)},
        {"Показатель": "Уникальных сессий", "Значение": len({e.get("sessionId") for e in events if e.get("sessionId")})},
    ]
    for name, count in sorted(counter.items()):
        rows.append({"Показатель": f"Событие: {name}", "Значение": count})
    for name, count in sorted(turn.items()):
        rows.append({"Показатель": f"turn_end → {name}", "Значение": count})
    for name, count in sorted(manual_failed.items()):
        rows.append({"Показатель": f"manual: не прошло {name}", "Значение": count})
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description="JSONL EOS log → Excel")
    parser.add_argument(
        "input",
        nargs="?",
        default="logs/eos-debug.jsonl",
        help="Путь к .jsonl (по умолчанию logs/eos-debug.jsonl)",
    )
    parser.add_argument(
        "-o",
        "--output",
        help="Выходной .xlsx (по умолчанию рядом с входом, суффикс -report.xlsx)",
    )
    args = parser.parse_args()

    in_path = Path(args.input)
    if not in_path.is_file():
        print(f"Файл не найден: {in_path}", file=sys.stderr)
        sys.exit(1)

    out_path = Path(args.output) if args.output else in_path.with_name(in_path.stem + "-report.xlsx")

    events = load_jsonl(in_path)
    if not events:
        print("Лог пуст.", file=sys.stderr)
        sys.exit(1)

    roles = session_roles(events)
    turn_all = [flatten_turn_end(ev, roles) for ev in events if ev.get("event") == "turn_end"]
    turn_manual = [r for r in turn_all if r.get("finishReason") == "manual"]
    turn_auto = [r for r in turn_all if r.get("finishReason") == "auto"]
    finals_rows: list[dict] = []
    for ev in events:
        if ev.get("event") == "turn_end":
            finals_rows.extend(flatten_finals_from_turn(ev))
    summary = build_summary(events, roles)

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Сводка"
    write_sheet(ws_sum, [("Показатель", "Показатель"), ("Значение", "Значение")], summary)

    ws_manual = wb.create_sheet("manual (проблема)")
    write_sheet(ws_manual, TURN_END_COLUMNS, turn_manual)

    ws_all = wb.create_sheet("все реплики")
    write_sheet(ws_all, TURN_END_COLUMNS, turn_all)

    ws_auto = wb.create_sheet("auto (ок)")
    write_sheet(ws_auto, TURN_END_COLUMNS, turn_auto)

    if finals_rows:
        ws_f = wb.create_sheet("finals (шаги SM)")
        write_sheet(ws_f, FINALS_COLUMNS, finals_rows)

    wb.save(out_path)
    print(f"Готово: {out_path}")
    print(f"  manual: {len(turn_manual)} | все: {len(turn_all)} | auto: {len(turn_auto)} | finals: {len(finals_rows)}")
    print("  Подсказка: logs/eos-turns.csv обновляется сервером автоматически — можно открыть напрямую в Excel.")


if __name__ == "__main__":
    main()
